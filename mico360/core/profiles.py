"""Company Profiles: branding + document-layout settings used when exporting
meeting minutes (letterhead, logo placement, footer, page numbering).

Profiles are stored as one JSON file each; logos are copied into the profiles
directory so a profile is self-contained and portable. Supports import/export
to JSON, CSV and XLSX.
"""
from __future__ import annotations

import csv
import json
import logging
import shutil
import uuid
from dataclasses import asdict, dataclass, field, fields
from pathlib import Path

from ..config import PROFILES_DIR

log = logging.getLogger("mico360.profiles")

LOGO_POSITIONS = ["left", "center", "right"]
PAGENUM_POSITIONS = [
    "footer-left", "footer-center", "footer-right",
    "header-left", "header-center", "header-right",
]
ALIGNMENTS = ["left", "center", "right"]


@dataclass
class CompanyProfile:
    id: str = ""
    name: str = "New Company"
    address: str = ""
    phone: str = ""
    email: str = ""
    website: str = ""
    logo_path: str = ""              # absolute path to copied logo
    logo_position: str = "left"      # left | center | right
    logo_width_mm: float = 35.0      # printed width; height auto-scales (quality kept)
    footer_text: str = ""
    footer_alignment: str = "center"
    show_page_numbers: bool = True
    page_number_position: str = "footer-right"
    page_number_format: str = "Page {n} of {total}"
    accent_color: str = "#2C7BE5"

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "CompanyProfile":
        valid = {f.name for f in fields(cls)}
        return cls(**{k: v for k, v in d.items() if k in valid})


class ProfileStore:
    def __init__(self, directory: Path = PROFILES_DIR):
        self.dir = directory
        self.logos = directory / "logos"
        self.dir.mkdir(parents=True, exist_ok=True)
        self.logos.mkdir(parents=True, exist_ok=True)

    # -- CRUD ---------------------------------------------------------------
    def list(self) -> list[CompanyProfile]:
        out = []
        for f in sorted(self.dir.glob("*.json")):
            try:
                out.append(CompanyProfile.from_dict(json.loads(f.read_text(encoding="utf-8"))))
            except Exception:
                log.warning("bad profile %s", f, exc_info=True)
        out.sort(key=lambda p: p.name.lower())
        return out

    def get(self, profile_id: str) -> CompanyProfile | None:
        f = self.dir / f"{profile_id}.json"
        if f.exists():
            return CompanyProfile.from_dict(json.loads(f.read_text(encoding="utf-8")))
        return None

    def save(self, profile: CompanyProfile) -> CompanyProfile:
        if not (profile.name or "").strip():
            profile.name = "Company"            # never persist a blank company name
        if not profile.id:
            profile.id = uuid.uuid4().hex[:10]
        (self.dir / f"{profile.id}.json").write_text(
            json.dumps(profile.to_dict(), indent=2), encoding="utf-8"
        )
        return profile

    def delete(self, profile_id: str) -> bool:
        f = self.dir / f"{profile_id}.json"
        if f.exists():
            p = self.get(profile_id)
            if p and p.logo_path and Path(p.logo_path).parent == self.logos:
                Path(p.logo_path).unlink(missing_ok=True)
            f.unlink()
            return True
        return False

    def set_logo(self, profile: CompanyProfile, source_image: str | Path) -> CompanyProfile:
        """Copy a logo into the profile store (kept full-resolution for quality)."""
        src = Path(source_image)
        if not src.exists():
            raise FileNotFoundError(src)
        pid = profile.id or uuid.uuid4().hex[:10]
        profile.id = pid
        dest = self.logos / f"{pid}{src.suffix.lower()}"
        shutil.copy2(src, dest)
        profile.logo_path = str(dest)
        return profile

    # -- Import -------------------------------------------------------------
    def import_file(self, path: str | Path) -> list[CompanyProfile]:
        path = Path(path)
        ext = path.suffix.lower()
        if ext == ".json":
            data = json.loads(path.read_text(encoding="utf-8"))
            records = data if isinstance(data, list) else [data]
        elif ext == ".csv":
            with path.open(newline="", encoding="utf-8-sig") as fh:
                records = list(csv.DictReader(fh))
        elif ext in (".xlsx", ".xlsm"):
            records = _read_xlsx(path)
        else:
            raise ValueError(f"Unsupported import format: {ext}")

        imported: list[CompanyProfile] = []
        for rec in records:
            rec = {k: _coerce(k, v) for k, v in rec.items()}
            prof = CompanyProfile.from_dict(rec)
            prof.id = ""  # always assign a fresh id on import
            imported.append(self.save(prof))
        log.info("imported %d profile(s) from %s", len(imported), path.name)
        return imported

    # -- Export -------------------------------------------------------------
    def export_file(self, profiles: list[CompanyProfile], path: str | Path) -> Path:
        path = Path(path)
        ext = path.suffix.lower()
        rows = [p.to_dict() for p in profiles]
        if ext == ".json":
            path.write_text(json.dumps(rows, indent=2), encoding="utf-8")
        elif ext == ".csv":
            cols = [f.name for f in fields(CompanyProfile)]
            with path.open("w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=cols)
                w.writeheader()
                w.writerows(rows)
        elif ext == ".xlsx":
            _write_xlsx(rows, path)
        else:
            raise ValueError(f"Unsupported export format: {ext}")
        return path


# -- helpers ----------------------------------------------------------------
_BOOL_FIELDS = {"show_page_numbers"}
_FLOAT_FIELDS = {"logo_width_mm"}


def _coerce(key: str, value):
    if value is None:
        return ""
    if key in _BOOL_FIELDS:
        return str(value).strip().lower() in ("1", "true", "yes", "y")
    if key in _FLOAT_FIELDS:
        try:
            return float(value)
        except (TypeError, ValueError):
            return 35.0
    return value


def _read_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()        # read_only workbooks hold the file open until closed
    if not rows:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        out.append({headers[i]: r[i] for i in range(len(headers)) if i < len(r)})
    return out


def _write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Company Profiles"
    cols = [f.name for f in fields(CompanyProfile)]
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    wb.save(path)
