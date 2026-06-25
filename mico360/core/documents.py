"""Extract text from uploaded documents so they can feed the transcript box.

Supports .txt/.md, .docx, and .pdf (best-effort). Images are accepted but only
OCR'd when an OCR backend is available; otherwise the user is told clearly.
"""
from __future__ import annotations

import logging
from pathlib import Path

log = logging.getLogger("mico360.documents")

TEXT_EXTS = {".txt", ".md", ".log", ".csv"}
DOC_EXTS = {".docx", ".pdf"} | TEXT_EXTS
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tiff", ".webp"}


def is_document(path: str | Path) -> bool:
    return Path(path).suffix.lower() in DOC_EXTS


def is_image(path: str | Path) -> bool:
    return Path(path).suffix.lower() in IMAGE_EXTS


def extract_text(path: str | Path) -> str:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in TEXT_EXTS:
        return p.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        return _from_docx(p)
    if ext == ".pdf":
        return _from_pdf(p)
    if ext in IMAGE_EXTS:
        return _from_image(p)
    raise ValueError(f"Unsupported document type: {ext}")


def _from_docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    parts = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(t for t in parts if t.strip())


def _from_pdf(p: Path) -> str:
    # Try a couple of optional backends; degrade gracefully.
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(str(p))
        return "\n".join((pg.extract_text() or "") for pg in reader.pages).strip()
    except Exception:
        pass
    try:
        from pdfminer.high_level import extract_text as _pm  # type: ignore
        return (_pm(str(p)) or "").strip()
    except Exception:
        raise RuntimeError(
            "PDF text extraction needs the optional 'pypdf' package. "
            "Install it with:  pip install pypdf"
        )


def _from_image(p: Path) -> str:
    try:
        import pytesseract  # type: ignore
        from PIL import Image
        return pytesseract.image_to_string(Image.open(p)).strip()
    except Exception:
        raise RuntimeError(
            "Reading text from images needs Tesseract OCR + the 'pytesseract' "
            "package. The image was attached but no text could be extracted."
        )
