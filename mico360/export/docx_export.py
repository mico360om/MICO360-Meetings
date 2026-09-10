"""Word (.docx) export with company letterhead, logo, footer and page numbers."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt, RGBColor

from ..core.profiles import CompanyProfile
from . import md_blocks

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _hex_rgb(color: str) -> RGBColor:
    color = (color or "#8B1E1E").lstrip("#")
    try:
        return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))
    except Exception:
        return RGBColor(0x2C, 0x7B, 0xE5)


def _add_field(paragraph, instr: str) -> None:
    """Insert a Word field (e.g. PAGE / NUMPAGES) into a paragraph."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText"); instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin); run._r.append(instr_el); run._r.append(fld_end)


def _build_header(doc, profile: CompanyProfile) -> None:
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = _ALIGN.get(profile.logo_position, WD_ALIGN_PARAGRAPH.LEFT)

    if profile.logo_path and Path(profile.logo_path).exists():
        try:
            p.add_run().add_picture(profile.logo_path, width=Mm(profile.logo_width_mm))
        except Exception:
            pass

    info = header.add_paragraph()
    info.alignment = _ALIGN.get(profile.logo_position, WD_ALIGN_PARAGRAPH.LEFT)
    name_run = info.add_run(profile.name)
    name_run.bold = True
    name_run.font.size = Pt(13)
    name_run.font.color.rgb = _hex_rgb(profile.accent_color)
    contact_bits = [b for b in (profile.address, profile.phone, profile.email, profile.website) if b]
    if contact_bits:
        cr = info.add_run("\n" + "  •  ".join(contact_bits))
        cr.font.size = Pt(8)
        cr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _build_footer(doc, profile: CompanyProfile) -> None:
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = _ALIGN.get(profile.footer_alignment, WD_ALIGN_PARAGRAPH.CENTER)
    if profile.footer_text:
        r = p.add_run(profile.footer_text)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    if profile.show_page_numbers:
        pn = footer.add_paragraph()
        pos = profile.page_number_position
        pn.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if pos.endswith("right")
                        else WD_ALIGN_PARAGRAPH.LEFT if pos.endswith("left")
                        else WD_ALIGN_PARAGRAPH.CENTER)
        fmt = profile.page_number_format or "Page {n} of {total}"
        # split on tokens, emitting fields for {n} and {total}
        for token in re.split(r"(\{n\}|\{total\})", fmt):
            if token == "{n}":
                _add_field(pn, "PAGE")
            elif token == "{total}":
                _add_field(pn, "NUMPAGES")
            elif token:
                run = pn.add_run(token)
                run.font.size = Pt(8)


def export_docx(minutes_md: str, path: str | Path,
                profile: CompanyProfile | None = None) -> Path:
    path = Path(path)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if profile:
        _build_header(doc, profile)
        _build_footer(doc, profile)

    for blk in md_blocks.parse(minutes_md):
        if blk.kind == "h1":
            h = doc.add_heading(level=0)
            run = h.add_run(blk.text)
            run.font.color.rgb = _hex_rgb(profile.accent_color if profile else "#8B1E1E")
        elif blk.kind == "h2":
            doc.add_heading(blk.text, level=1)
        elif blk.kind == "kv":
            p = doc.add_paragraph()
            p.add_run(f"{blk.key}: ").bold = True
            p.add_run(blk.text)
        elif blk.kind == "bullet":
            for it in blk.items:
                p = doc.add_paragraph(style="List Bullet")
                for txt, bold in md_blocks.runs(it):
                    p.add_run(txt).bold = bold
        elif blk.kind == "table":
            cols = len(blk.headers)
            table = doc.add_table(rows=1, cols=cols)
            table.style = "Light Grid Accent 1"
            for j, htext in enumerate(blk.headers):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].add_run(htext).bold = True
            for row in blk.rows:
                cells = table.add_row().cells
                for j in range(cols):
                    cells[j].text = row[j] if j < len(row) else ""
        elif blk.kind == "para":
            p = doc.add_paragraph()
            for txt, bold in md_blocks.runs(blk.text):
                p.add_run(txt).bold = bold

    doc.save(str(path))
    return path
